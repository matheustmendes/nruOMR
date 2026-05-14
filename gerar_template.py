import sys
import yaml
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth

# --- CONFIGURAÇÕES DE LAYOUT ---
MARCADOR_TAM = 5 * mm
MARCADOR_MARGEM = 15 * mm

CIRCULO_RAIO = 2.5 * mm
CIRCULO_BORDA = 1.5
CIRCULO_ESPACO_AJ = 8 * mm
CIRCULO_ESPACO_DIA = 9 * mm

LINHA_ALTURA = 8 * mm

FONTE_TITULO = ("Helvetica-Bold", 13)
FONTE_SUBTITULO = ("Helvetica", 11)
FONTE_CABECALHO = ("Helvetica-Bold", 9)
FONTE_CORPO = ("Helvetica", 11)

MARGEM_ESQUERDA = 15 * mm
MARGEM_DIREITA = 15 * mm

# Mapeamento de abas para configuração de dias
CONFIG_ABAS = {
    "ONDINA IMPRESSÃO":      ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado"],
    "CANELA IMPRESSÃO":      ["Segunda", "Terça", "Quarta", "Quinta", "Sexta"],
    "SÃO LÁZARO IMPRESSÃO":  ["Segunda", "Terça", "Quarta", "Quinta", "Sexta"],
    " PDCA FDS":             ["Sábado"],
    "PDSL FDS":              ["Sábado"],
}


# --- LEITURA DA PLANILHA ---

def ler_planilha(caminho_xlsx: str, nome_aba: str) -> dict:
    wb = load_workbook(caminho_xlsx, read_only=True)

    if nome_aba not in wb.sheetnames:
        raise ValueError(
            f"Aba '{nome_aba}' não encontrada. "
            f"Abas disponíveis: {wb.sheetnames}"
        )

    ws = wb[nome_aba]

    linhas_cabecalho = []
    for row in ws.iter_rows(max_row=5, values_only=True):
        valor = None
        for cell in row:
            if cell and str(cell).strip():
                valor = str(cell).strip()
                break
        linhas_cabecalho.append(valor)

    restaurante = ""
    if linhas_cabecalho[2]:
        restaurante = linhas_cabecalho[2].replace("RELAÇÃO DE BOLSISTAS", "").strip()

    mes_ano = linhas_cabecalho[3] or ""
    datas = linhas_cabecalho[4] or ""

    alunos = []
    for row in ws.iter_rows(min_row=8, values_only=True):
        if len(row) < 3:
            continue

        nome = row[1]
        matricula = row[2]

        if not nome or not str(nome).strip():
            continue

        nome = str(nome).strip()

        if matricula is None:
            matricula = ""
        elif isinstance(matricula, float):
            matricula = str(int(matricula))
        else:
            matricula = str(matricula).strip()

        alunos.append((nome, matricula))

    wb.close()

    return {
        "restaurante": restaurante,
        "mes_ano": mes_ano,
        "datas": datas,
        "alunos": alunos,
    }


# --- ABREVIAÇÃO DE NOMES ---

def abreviar_nome(nome: str, largura_max_pt: float, fonte: str, tamanho: float) -> str:
    """
    Abrevia sobrenomes do final para o início até o nome caber na largura disponível.

    Estratégia:
        1. Se o nome já cabe, retorna sem alteração.
        2. Caso contrário, substitui sobrenomes (do último para o primeiro)
           pela inicial seguida de ponto: "Silva" → "S."
        3. Preserva sempre o primeiro nome e o segundo nome inteiros.
           Só abrevia a partir do terceiro token em diante (de trás pra frente).
        4. Se mesmo abreviando tudo ainda não couber, trunca com "…".

    Exemplos:
        "Paulo Matheus Silva Santos Marques"
            → "Paulo Matheus S. S. M."  (se necessário)
        "Ana Beatriz"
            → "Ana Beatriz"  (não precisa abreviar)
        "Ana Beatriz Carvalho"
            → "Ana Beatriz C."  (se necessário)
    """
    # Já cabe? Retorna sem alterar.
    if stringWidth(nome, fonte, tamanho) <= largura_max_pt:
        return nome

    partes = nome.split()

    # Menos de 3 partes: não tem o que abreviar, trunca se necessário.
    if len(partes) < 3:
        return _truncar(nome, largura_max_pt, fonte, tamanho)

    # Abrevia do último sobrenome em direção ao terceiro token
    # (preserva partes[0] e partes[1] sempre inteiros)
    abreviadas = partes[:]
    for i in range(len(partes) - 1, 1, -1):
        abreviadas[i] = partes[i][0].upper() + "."
        candidato = " ".join(abreviadas)
        if stringWidth(candidato, fonte, tamanho) <= largura_max_pt:
            return candidato

    # Ainda não coube mesmo abreviando tudo — trunca
    return _truncar(" ".join(abreviadas), largura_max_pt, fonte, tamanho)


def _truncar(nome: str, largura_max_pt: float, fonte: str, tamanho: float) -> str:
    """Trunca com reticências se o nome não couber de jeito nenhum."""
    sufixo = "…"
    while nome and stringWidth(nome + sufixo, fonte, tamanho) > largura_max_pt:
        nome = nome[:-1]
    return nome + sufixo


# --- DESENHO DO PDF ---

def desenhar_marcadores(c, largura, altura):
    """Desenha os 4 marcadores fiduciais nos cantos."""
    c.setFillColor("black")
    c.setStrokeColor("black")

    posicoes = [
        (MARCADOR_MARGEM, altura - MARCADOR_MARGEM - MARCADOR_TAM),
        (largura - MARCADOR_MARGEM - MARCADOR_TAM, altura - MARCADOR_MARGEM - MARCADOR_TAM),
        (MARCADOR_MARGEM, MARCADOR_MARGEM),
        (largura - MARCADOR_MARGEM - MARCADOR_TAM, MARCADOR_MARGEM),
    ]

    for (x, y) in posicoes:
        c.rect(x, y, MARCADOR_TAM, MARCADOR_TAM, fill=1)


def desenhar_cabecalho(c, largura, altura, pagina_atual, total_paginas, info):
    """Desenha o cabeçalho institucional com dados da planilha."""
    y = altura - MARCADOR_MARGEM - MARCADOR_TAM - 12 * mm

    c.setFont(*FONTE_TITULO)
    c.setFillColor("black")
    c.drawCentredString(largura / 2, y, "Pró-Reitoria de Assistência Estudantil")

    y -= 5 * mm
    c.drawCentredString(largura / 2, y, f"RELAÇÃO DE BOLSISTAS {info['restaurante']}")

    y -= 5 * mm
    c.setFont(*FONTE_SUBTITULO)
    c.drawCentredString(largura / 2, y, f"{info['mes_ano']}  |  {info['datas']}")

    y -= 5 * mm
    c.setFont(*FONTE_CORPO)
    c.drawRightString(largura - MARGEM_DIREITA, y, f"Página {pagina_atual} de {total_paginas}")

    return y


def calcular_layout_colunas(posicoes_circulos):
    """
    Calcula as posições das colunas Nome e Matrícula em função dos círculos.

    A matrícula é alinhada à DIREITA com 2mm de folga antes do primeiro círculo.
    O espaço do nome é calculado com base na largura real (pior caso) da matrícula.
    """
    nome_x = MARGEM_ESQUERDA + 8 * mm

    primeiro_circulo_esquerda = posicoes_circulos[0][0] - CIRCULO_RAIO
    matricula_right = primeiro_circulo_esquerda - 2 * mm

    max_matricula_pt = stringWidth("222117117", FONTE_CORPO[0], FONTE_CORPO[1])

    nome_right_max = matricula_right - max_matricula_pt - 2 * mm
    nome_largura_max_pt = nome_right_max - nome_x

    return {
        "nome_x": nome_x,
        "nome_largura_max_pt": nome_largura_max_pt,
        "matricula_right": matricula_right,
    }


def calcular_posicoes_circulos(largura, dias):
    """
    Calcula as posições X de cada círculo.
    Retorna uma lista de (x_centro_A, x_centro_J) por dia.
    """
    margem_direita = 15 * mm
    margem_seguranca = 5 * mm

    n = len(dias)
    espaco_total = (n - 1) * (CIRCULO_ESPACO_AJ + CIRCULO_ESPACO_DIA) + CIRCULO_ESPACO_AJ + CIRCULO_RAIO

    x_inicio_max = largura - margem_direita - margem_seguranca - espaco_total
    x_inicio_padrao = 105 * mm

    x_inicio = min(x_inicio_padrao, x_inicio_max)

    posicoes = []
    for i in range(n):
        x_a = x_inicio + i * (CIRCULO_ESPACO_AJ + CIRCULO_ESPACO_DIA)
        x_j = x_a + CIRCULO_ESPACO_AJ
        posicoes.append((x_a, x_j))

    return posicoes


def desenhar_cabecalho_tabela(c, y, posicoes_circulos, largura, dias, layout):
    """Desenha o cabeçalho da tabela."""
    c.setFont(*FONTE_CABECALHO)
    c.setFillColor("black")

    c.drawString(MARGEM_ESQUERDA, y, "Nº")
    c.drawString(layout["nome_x"], y, "Nome")
    c.drawRightString(layout["matricula_right"], y, "Matrícula")

    for i, dia in enumerate(dias):
        x_a, x_j = posicoes_circulos[i]
        x_centro = (x_a + x_j) / 2
        c.drawCentredString(x_centro, y, dia)

    y -= 4 * mm
    c.setFont(*FONTE_CORPO)
    for i in range(len(dias)):
        x_a, x_j = posicoes_circulos[i]
        c.drawCentredString(x_a, y, "A")
        c.drawCentredString(x_j, y, "J")

    y -= 2 * mm
    c.setStrokeColor("black")
    c.setLineWidth(0.5)
    c.line(MARGEM_ESQUERDA, y, largura - MARGEM_DIREITA, y)

    return y


def desenhar_linha_aluno(c, y, numero, nome, matricula, posicoes_circulos, layout):
    """Desenha uma linha completa: número, nome, matrícula e círculos."""
    c.setFillColor("black")
    c.setFont(*FONTE_CORPO)

    # Abrevia sobrenomes se necessário — fonte sempre fixa
    nome_renderizado = abreviar_nome(
        nome,
        layout["nome_largura_max_pt"],
        FONTE_CORPO[0],
        FONTE_CORPO[1],
    )

    c.drawRightString(MARGEM_ESQUERDA + 6 * mm, y, str(numero))
    c.drawString(layout["nome_x"], y, nome_renderizado)
    c.drawRightString(layout["matricula_right"], y, matricula)

    c.setStrokeColor("black")
    c.setLineWidth(CIRCULO_BORDA)
    c.setFillColor("white")

    y_circulo = y + 1.2 * mm

    for x_a, x_j in posicoes_circulos:
        c.circle(x_a, y_circulo, CIRCULO_RAIO, fill=1)
        c.circle(x_j, y_circulo, CIRCULO_RAIO, fill=1)


def desenhar_rodape(c, largura):
    """Desenha instrução no rodapé."""
    c.setFont(*FONTE_CORPO)
    c.setFillColor("black")
    y = MARCADOR_MARGEM + MARCADOR_TAM + 4 * mm
    c.drawCentredString(
        largura / 2, y,
        "INSTRUÇÃO: Preencha completamente o círculo ( ● ) para registrar presença."
    )


def gerar_template(info: dict, dias: list, arquivo_saida: str):
    """Gera o PDF completo do template."""
    alunos = info["alunos"]

    c = canvas.Canvas(arquivo_saida, pagesize=A4)
    largura, altura = A4

    posicoes_circulos = calcular_posicoes_circulos(largura, dias)
    layout = calcular_layout_colunas(posicoes_circulos)

    y_limite_inferior = MARCADOR_MARGEM + MARCADOR_TAM + 10 * mm
    y_inicio_tabela = altura - MARCADOR_MARGEM - MARCADOR_TAM - 12 * mm - 15 * mm - 6 * mm - 6 * mm - 3 * mm
    espaco_disponivel = y_inicio_tabela - y_limite_inferior
    alunos_por_pagina = int(espaco_disponivel / LINHA_ALTURA)

    total_paginas = (len(alunos) + alunos_por_pagina - 1) // alunos_por_pagina
    pagina_atual = 1

    for idx_inicio in range(0, len(alunos), alunos_por_pagina):
        alunos_pagina = alunos[idx_inicio:idx_inicio + alunos_por_pagina]

        desenhar_marcadores(c, largura, altura)
        y = desenhar_cabecalho(c, largura, altura, pagina_atual, total_paginas, info)

        y -= 6 * mm
        y = desenhar_cabecalho_tabela(c, y, posicoes_circulos, largura, dias, layout)

        y -= 5 * mm
        y_topo_tabela = y + 2 * mm

        if pagina_atual == 1:
            primeiro_circulo_y = y + 1.2 * mm

        for i, (nome, matricula) in enumerate(alunos_pagina):
            numero = idx_inicio + i + 1
            desenhar_linha_aluno(c, y, numero, nome, matricula, posicoes_circulos, layout)
            y -= LINHA_ALTURA
        y_base_tabela = y + LINHA_ALTURA - 4 * mm

        c.setStrokeColor("black")
        c.setLineWidth(0.3)
        for i in range(1, len(dias)):
            x_a_atual = posicoes_circulos[i][0]
            x_j_anterior = posicoes_circulos[i - 1][1]
            x_linha = (x_j_anterior + x_a_atual) / 2
            c.line(x_linha, y_topo_tabela, x_linha, y_base_tabela)

        desenhar_rodape(c, largura)

        if idx_inicio + alunos_por_pagina < len(alunos):
            c.showPage()
            pagina_atual += 1

    c.save()

    print(f"Template gerado: {arquivo_saida}")
    print(f"  Restaurante: {info['restaurante']}")
    print(f"  {info['mes_ano']}  |  {info['datas']}")
    print(f"  {len(alunos)} alunos em {total_paginas} página(s)")
    print(f"  {alunos_por_pagina} alunos por página")
    print(f"  {len(dias)} dias × 2 (A/J) = {len(dias) * 2} círculos por linha")

    return {
        "posicoes_circulos": posicoes_circulos,
        "alunos_por_pagina": alunos_por_pagina,
        "total_paginas": total_paginas,
        "primeiro_circulo_y_reportlab": primeiro_circulo_y,
    }


def gerar_config(info: dict, dias: list, resultado_template: dict, arquivo_config: str):
    largura_mm, altura_mm = A4[0] / mm, A4[1] / mm
    posicoes = resultado_template["posicoes_circulos"]

    marcador_mm = MARCADOR_TAM / mm
    margem_mm = MARCADOR_MARGEM / mm

    marcadores = {
        "superior_esquerdo": {"x": margem_mm, "y": margem_mm},
        "superior_direito":  {"x": largura_mm - margem_mm - marcador_mm, "y": margem_mm},
        "inferior_esquerdo": {"x": margem_mm, "y": altura_mm - margem_mm - marcador_mm},
        "inferior_direito":  {"x": largura_mm - margem_mm - marcador_mm, "y": altura_mm - margem_mm - marcador_mm},
    }

    circulos_por_dia = {}
    for i, dia in enumerate(dias):
        x_a, x_j = posicoes[i]
        circulos_por_dia[dia] = {
            "almoco_x": round(x_a / mm, 1),
            "janta_x": round(x_j / mm, 1),
        }

    config = {
        "pagina": {
            "largura_mm": round(largura_mm, 1),
            "altura_mm": round(altura_mm, 1),
            "orientacao": "portrait",
        },
        "scan": {
            "dpi": 200,
        },
        "marcadores": {
            "tamanho_mm": round(marcador_mm, 1),
            "margem_mm": round(margem_mm, 1),
            "posicoes": marcadores,
        },
        "circulos": {
            "raio_mm": round(CIRCULO_RAIO / mm, 1),
            "borda_pt": CIRCULO_BORDA,
        },
        "layout": {
            "alunos_por_pagina": resultado_template["alunos_por_pagina"],
            "linha_altura_mm": round(LINHA_ALTURA / mm, 1),
            "primeira_linha_y_mm": round(
                altura_mm - resultado_template["primeiro_circulo_y_reportlab"] / mm,
                1
            ),
            "dias": dias,
            "circulos_por_dia": circulos_por_dia,
        },
        "restaurante": info["restaurante"],
    }

    with open(arquivo_config, "w", encoding="utf-8") as f:
        yaml.dump(config, f, default_flow_style=False, allow_unicode=True, sort_keys=False)

    print(f"Config gerado: {arquivo_config}")


# --- EXECUÇÃO ---

def main():
    if len(sys.argv) < 3:
        print("Uso: python gerar_template.py <planilha.xlsx> <nome_da_aba>")
        print()
        print("Exemplo:")
        print('  python gerar_template.py alunos.xlsx "CANELA IMPRESSÃO"')
        print()

        if len(sys.argv) == 2:
            wb = load_workbook(sys.argv[1], read_only=True)
            print(f"Abas disponíveis em {sys.argv[1]}:")
            for nome in wb.sheetnames:
                print(f"  - {nome}")
            wb.close()

        sys.exit(1)

    caminho_xlsx = sys.argv[1]
    nome_aba = sys.argv[2]

    info = ler_planilha(caminho_xlsx, nome_aba)
    print(f"Lidos {len(info['alunos'])} alunos da aba '{nome_aba}'")

    if nome_aba in CONFIG_ABAS:
        dias = CONFIG_ABAS[nome_aba]
    else:
        print(f"Aviso: aba '{nome_aba}' não tem configuração de dias definida.")
        print(f"Usando dias padrão (Seg-Sex).")
        dias = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta"]

    nome_restaurante = info["restaurante"].lower().replace(" ", "_") or "template"
    arquivo_saida = f"template_{nome_restaurante}.pdf"
    arquivo_config = f"config_{nome_restaurante}.yaml"

    resultado = gerar_template(info, dias, arquivo_saida)
    gerar_config(info, dias, resultado, arquivo_config)

    print(f"\nPosições dos círculos (em mm):")
    for i, dia in enumerate(dias):
        x_a, x_j = resultado["posicoes_circulos"][i]
        print(f"  {dia}: A = {x_a/mm:.1f}mm, J = {x_j/mm:.1f}mm")


if __name__ == "__main__":
    main()