"""
importar_legado.py

Importa frequências históricas para a planilha de presenças do restaurante,
no mesmo formato gerado pelo processo automatizado.

Formato do xlsx (linha 5 = cabeçalho):
  A: Nome completo
  B: Matrícula
  C, D, E...: frequência de cada período (valor = total de dias presentes)
              cabeçalho na linha 5 = período, ex: "04/05 - 09/05"

Uso:
  python importar_legado.py planilha.xlsx ondina
  python importar_legado.py planilha.xlsx canela --forcar
"""

import sys
import os
import re

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import load_workbook
from google_sheets import exportar_para_sheets

_SEP = re.compile(r"\s+[–—\-à]\s+")


def _normalizar(txt):
    return _SEP.sub(" a ", str(txt).strip())


def ler_xlsx(caminho):
    wb = load_workbook(caminho, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 6:
        raise ValueError("Arquivo deve ter pelo menos 6 linhas (linha 5 = cabeçalho de períodos).")

    header = rows[4]

    periodos = []
    for j in range(2, len(header)):
        val = header[j]
        if val and str(val).strip():
            txt = _normalizar(str(val))
            if re.search(r"\d{1,2}/\d{1,2}", txt):
                periodos.append((j, txt))

    if not periodos:
        raise ValueError("Nenhum período encontrado na linha 5 (colunas C em diante).")

    alunos = []
    freqs = {p: [] for _, p in periodos}

    for row in rows[5:]:
        if not row or not any(c is not None for c in row):
            continue
        nome = row[0]
        if not nome or not str(nome).strip():
            continue
        nome = str(nome).strip()

        mat = row[1] if len(row) > 1 else None
        if mat is None:
            mat = ""
        elif isinstance(mat, float):
            mat = str(int(mat))
        else:
            mat = str(mat).strip()

        alunos.append((nome, mat))

        for j, periodo in periodos:
            val = row[j] if j < len(row) else None
            try:
                freq = int(float(str(val))) if val is not None and str(val).strip() else 0
            except (ValueError, TypeError):
                freq = 0
            freqs[periodo].append(freq)

    return alunos, [(p, freqs[p]) for _, p in periodos]


def main():
    if len(sys.argv) < 3:
        print("Uso: python importar_legado.py <planilha.xlsx> <restaurante> [--forcar]")
        print("Restaurantes: canela, ondina, sao_lazaro")
        sys.exit(1)

    caminho = sys.argv[1]
    restaurante_key = sys.argv[2]
    forcar = "--forcar" in sys.argv

    if not os.path.isfile(caminho):
        print(f"Arquivo nao encontrado: {caminho}")
        sys.exit(1)

    alunos, semanas = ler_xlsx(caminho)

    print(f"Restaurante : {restaurante_key}")
    print(f"Alunos      : {len(alunos)}")
    print(f"Semanas     : {[p for p, _ in semanas]}")
    print()

    for periodo, freq_alunos in semanas:
        contagem = [
            {"numero": i + 1, "presencas": freq, "detalhes": {}}
            for i, freq in enumerate(freq_alunos)
        ]

        resultado = exportar_para_sheets(
            contagem, alunos, dias=[],
            restaurante_key=restaurante_key,
            periodo=periodo,
            forcar=forcar,
        )

        if resultado.get("ok"):
            print(f"[OK]        {periodo} -> aba {resultado['aba']}")
        elif resultado.get("duplicado"):
            print(f"[DUPLICADO] {periodo} ja existe em {resultado['aba']} -- use --forcar para substituir")
        else:
            print(f"[ERRO]      {periodo}: {resultado.get('erro')}")


if __name__ == "__main__":
    main()
