import sys
import yaml
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm

# --- CONFIGURAÇÕES DE LAYOUT ---
MARCADOR_TAM = 5 * mm
MARCADOR_MARGEM = 15 * mm

CIRCULO_RAIO = 2.5 * mm
CIRCULO_BORDA = 1.5
CIRCULO_ESPACO_AJ = 8 * mm
CIRCULO_ESPACO_DIA = 9 * mm

LINHA_ALTURA = 8 * mm

FONTE_TITULO = ("Helvetica-Bold", 11)
FONTE_SUBTITULO = ("Helvetica", 9)
FONTE_CABECALHO = ("Helvetica-Bold", 7)
FONTE_CORPO = ("Helvetica", 9)

MARGEM_ESQUERDA = 15 * mm
MARGEM_DIREITA = 15 * mm

# Mapeamento de abas para configuração de dias
CONFIG_ABAS = {
    "ONDINA IMPRESSÃO":      ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado"],
    "CANELA IMPRESSÃO":      ["Segunda", "Terça", "Quarta", "Quinta", "Sexta"],
    "SÃO LÁZARO IMPRESSÃO":  ["Segunda", "Terça", "Quarta", "Quinta", "Sexta"],
    " PDCA FDS":             ["Sábado"],
    "PDSL FDS":              ["Sábado"],
}


# --- LEITURA DA PLANILHA ---

def ler_planilha(caminho_xlsx: str, nome_aba: str) -> dict:
    """
    Lê uma aba da planilha e extrai cabeçalho + lista de alunos.

    A planilha tem essa estrutura fixa:
        Linha 1: vazia
        Linha 2: "Pró-Reitoria de Assistência Estudantil"
        Linha 3: "RELAÇÃO DE BOLSISTAS [RESTAURANTE]"
        Linha 4: "MÊS [MÊS] DE [ANO]"
        Linha 5: "DATA: [inicio] - [fim]"
        Linha 6: cabeçalho (Nº, NOME, Matrícula, dias...)
        Linha 7: sub-cabeçalho (A, J, A, J...)
        Linha 8+: dados dos alunos

    Returns:
        dict com chaves: restaurante, mes_ano, datas, alunos[(nome, matricula)]
    """
    wb = load_workbook(caminho_xlsx, read_only=True)

    if nome_aba not in wb.sheetnames:
        print(f"Erro: aba '{nome_aba}' não encontrada.")
        print(f"Abas disponíveis: {wb.sheetnames}")
        sys.exit(1)

    ws = wb[nome_aba]

    # Extrai cabeçalho das primeiras linhas
    linhas_cabecalho = []
    for row in ws.iter_rows(max_row=5, values_only=True):
        # Pega o primeiro valor não-None da linha
        valor = None
        for cell in row:
            if cell and str(cell).strip():
                valor = str(cell).strip()
                break
        linhas_cabecalho.append(valor)

    # Linha 3: "RELAÇÃO DE BOLSISTAS CANELA" -> extrai "CANELA"
    restaurante = ""
    if linhas_cabecalho[2]:
        restaurante = linhas_cabecalho[2].replace("RELAÇÃO DE BOLSISTAS", "").strip()

    # Linha 4: "MÊS ABRIL DE 2026"
    mes_ano = linhas_cabecalho[3] or ""

    # Linha 5: "DATA: 22/04 - 25/04"
    datas = linhas_cabecalho[4] or ""

    # Extrai alunos a partir da linha 8
    alunos = []
    for row in ws.iter_rows(min_row=8, values_only=True):
        if len(row) < 3:
            continue

        nome = row[1]
        matricula = row[2]

        # Pula linhas sem nome
        if not nome or not str(nome).strip():
            continue

        nome = str(nome).strip()

        # Normaliza matrícula: pode vir como float (222117117.0) ou string
        if matricula is None:
            matricula = ""
        elif isinstance(matricula, float):
            matricula = str(int(matricula))
        else:
            matricula = str(matricula).strip()

        alunos.append((nome, matricula))

    wb.close()

    return {
        "restaurante": restaurante,
        "mes_ano": mes_ano,
        "datas": datas,
        "alunos": alunos,
    }


# --- DESENHO DO PDF ---

def desenhar_marcadores(c, largura, altura):
    """Desenha os 4 marcadores fiduciais nos cantos."""
    c.setFillColor("black")
    c.setStrokeColor("black")

    posicoes = [
        (MARCADOR_MARGEM, altura - MARCADOR_MARGEM - MARCADOR_TAM),
        (largura - MARCADOR_MARGEM - MARCADOR_TAM, altura - MARCADOR_MARGEM - MARCADOR_TAM),
        (MARCADOR_MARGEM, MARCADOR_MARGEM),
        (largura - MARCADOR_MARGEM - MARCADOR_TAM, MARCADOR_MARGEM),
    ]

    for (x, y) in posicoes:
        c.rect(x, y, MARCADOR_TAM, MARCADOR_TAM, fill=1)


def desenhar_cabecalho(c, largura, altura, pagina_atual, total_paginas, info):
    """Desenha o cabeçalho institucional com dados da planilha."""
    y = altura - MARCADOR_MARGEM - MARCADOR_TAM - 12 * mm

    c.setFont(*FONTE_TITULO)
    c.setFillColor("black")
    c.drawCentredString(largura / 2, y, "Pró-Reitoria de Assistência Estudantil")

    y -= 5 * mm
    c.drawCentredString(largura / 2, y, f"RELAÇÃO DE BOLSISTAS {info['restaurante']}")

    y -= 5 * mm
    c.setFont(*FONTE_SUBTITULO)
    c.drawCentredString(largura / 2, y, f"{info['mes_ano']}  |  {info['datas']}")

    y -= 5 * mm
    c.setFont(*FONTE_CORPO)
    c.drawRightString(largura - MARGEM_DIREITA, y, f"Página {pagina_atual} de {total_paginas}")

    return y


def calcular_posicoes_circulos(largura, dias):
    """
    Calcula as posições X de cada círculo.
    Retorna uma lista de (x_centro_A, x_centro_J) por dia.

    O x_inicio é calculado dinamicamente pra garantir que todos
    os círculos caibam na página com margem de segurança.
    """
    margem_direita = 15 * mm
    margem_seguranca = 5 * mm  # folga extra da borda

    # Calcula o espaço total que os círculos ocupam
    n = len(dias)
    espaco_total = (n - 1) * (CIRCULO_ESPACO_AJ + CIRCULO_ESPACO_DIA) + CIRCULO_ESPACO_AJ + CIRCULO_RAIO

    # x_inicio = largura - margem - segurança - espaço total
    x_inicio_max = largura - margem_direita - margem_seguranca - espaco_total
    x_inicio_padrao = 105 * mm

    # Usa o padrão se cabe, senão recua
    x_inicio = min(x_inicio_padrao, x_inicio_max)

    posicoes = []
    for i in range(n):
        x_a = x_inicio + i * (CIRCULO_ESPACO_AJ + CIRCULO_ESPACO_DIA)
        x_j = x_a + CIRCULO_ESPACO_AJ
        posicoes.append((x_a, x_j))

    return posicoes


def desenhar_cabecalho_tabela(c, y, posicoes_circulos, largura, dias):
    """Desenha o cabeçalho da tabela."""
    c.setFont(*FONTE_CABECALHO)
    c.setFillColor("black")

    c.drawString(MARGEM_ESQUERDA, y, "Nº")
    c.drawString(MARGEM_ESQUERDA + 8 * mm, y, "Nome")
    c.drawString(78 * mm, y, "Matrícula")

    for i, dia in enumerate(dias):
        x_a, x_j = posicoes_circulos[i]
        x_centro = (x_a + x_j) / 2
        c.drawCentredString(x_centro, y, dia)

    y -= 4 * mm
    c.setFont(*FONTE_CORPO)
    for i in range(len(dias)):
        x_a, x_j = posicoes_circulos[i]
        c.drawCentredString(x_a, y, "A")
        c.drawCentredString(x_j, y, "J")

    y -= 2 * mm
    c.setStrokeColor("black")
    c.setLineWidth(0.5)
    c.line(MARGEM_ESQUERDA, y, largura - MARGEM_DIREITA, y)

    return y


def desenhar_linha_aluno(c, y, numero, nome, matricula, posicoes_circulos):
    """Desenha uma linha completa: número, nome, matrícula e círculos."""
    c.setFont(*FONTE_CORPO)
    c.setFillColor("black")

    nome_max = 38
    if len(nome) > nome_max:
        nome = nome[:nome_max - 2] + ".."

    c.drawRightString(MARGEM_ESQUERDA + 6 * mm, y, str(numero))
    c.drawString(MARGEM_ESQUERDA + 8 * mm, y, nome)
    c.drawString(78 * mm, y, matricula)

    c.setStrokeColor("black")
    c.setLineWidth(CIRCULO_BORDA)
    c.setFillColor("white")

    y_circulo = y + 1.2 * mm

    for x_a, x_j in posicoes_circulos:
        c.circle(x_a, y_circulo, CIRCULO_RAIO, fill=1)
        c.circle(x_j, y_circulo, CIRCULO_RAIO, fill=1)


def desenhar_rodape(c, largura):
    """Desenha instrução no rodapé."""
    c.setFont(*FONTE_CORPO)
    c.setFillColor("black")
    y = MARCADOR_MARGEM + MARCADOR_TAM + 4 * mm
    c.drawCentredString(
        largura / 2, y,
        "INSTRUÇÃO: Preencha completamente o círculo ( ● ) para registrar presença."
    )


def gerar_template(info: dict, dias: list, arquivo_saida: str):
    """Gera o PDF completo do template."""
    alunos = info["alunos"]

    c = canvas.Canvas(arquivo_saida, pagesize=A4)
    largura, altura = A4

    posicoes_circulos = calcular_posicoes_circulos(largura, dias)

    # Calcula quantos alunos cabem por página
    y_limite_inferior = MARCADOR_MARGEM + MARCADOR_TAM + 10 * mm
    y_inicio_tabela = altura - MARCADOR_MARGEM - MARCADOR_TAM - 12 * mm - 15 * mm - 6 * mm - 6 * mm - 3 * mm
    espaco_disponivel = y_inicio_tabela - y_limite_inferior
    alunos_por_pagina = int(espaco_disponivel / LINHA_ALTURA)

    total_paginas = (len(alunos) + alunos_por_pagina - 1) // alunos_por_pagina
    pagina_atual = 1

    for idx_inicio in range(0, len(alunos), alunos_por_pagina):
        alunos_pagina = alunos[idx_inicio:idx_inicio + alunos_por_pagina]

        desenhar_marcadores(c, largura, altura)

        y = desenhar_cabecalho(c, largura, altura, pagina_atual, total_paginas, info)

        y -= 6 * mm
        y = desenhar_cabecalho_tabela(c, y, posicoes_circulos, largura, dias)

        y -= 5 * mm
        y_topo_tabela = y + 2 * mm

        # Salva o Y do centro do primeiro círculo (só na primeira página)
        if pagina_atual == 1:
            primeiro_circulo_y = y + 1.2 * mm  # mesmo ajuste de desenhar_linha_aluno

        for i, (nome, matricula) in enumerate(alunos_pagina):
            numero = idx_inicio + i + 1
            desenhar_linha_aluno(c, y, numero, nome, matricula, posicoes_circulos)
            y -= LINHA_ALTURA
        y_base_tabela = y + LINHA_ALTURA - 4 * mm

        # Linhas verticais entre os dias
        c.setStrokeColor("black")
        c.setLineWidth(0.3)
        for i in range(1, len(dias)):
            x_a_atual = posicoes_circulos[i][0]
            x_j_anterior = posicoes_circulos[i - 1][1]
            x_linha = (x_j_anterior + x_a_atual) / 2
            c.line(x_linha, y_topo_tabela, x_linha, y_base_tabela)

        desenhar_rodape(c, largura)

        if idx_inicio + alunos_por_pagina < len(alunos):
            c.showPage()
            pagina_atual += 1

    c.save()

    print(f"Template gerado: {arquivo_saida}")
    print(f"  Restaurante: {info['restaurante']}")
    print(f"  {info['mes_ano']}  |  {info['datas']}")
    print(f"  {len(alunos)} alunos em {total_paginas} página(s)")
    print(f"  {alunos_por_pagina} alunos por página")
    print(f"  {len(dias)} dias × 2 (A/J) = {len(dias) * 2} círculos por linha")

    return {
        "posicoes_circulos": posicoes_circulos,
        "alunos_por_pagina": alunos_por_pagina,
        "total_paginas": total_paginas,
        "primeiro_circulo_y_reportlab": primeiro_circulo_y,
    }


def gerar_config(info: dict, dias: list, resultado_template: dict, arquivo_config: str):
    """
    Gera o config.yaml com todas as informações que o OMR precisa
    pra saber onde procurar os círculos no scan.

    As posições são salvas em mm — na hora de usar, o pipeline
    converte pra pixels baseado no DPI do scan.
    """
    largura_mm, altura_mm = A4[0] / mm, A4[1] / mm
    posicoes = resultado_template["posicoes_circulos"]

    # Posições dos marcadores fiduciais (em mm)
    marcador_mm = MARCADOR_TAM / mm
    margem_mm = MARCADOR_MARGEM / mm

    marcadores = {
        "superior_esquerdo": {"x": margem_mm, "y": margem_mm},
        "superior_direito":  {"x": largura_mm - margem_mm - marcador_mm, "y": margem_mm},
        "inferior_esquerdo": {"x": margem_mm, "y": altura_mm - margem_mm - marcador_mm},
        "inferior_direito":  {"x": largura_mm - margem_mm - marcador_mm, "y": altura_mm - margem_mm - marcador_mm},
    }

    # Posições dos círculos por dia (em mm)
    circulos_por_dia = {}
    for i, dia in enumerate(dias):
        x_a, x_j = posicoes[i]
        circulos_por_dia[dia] = {
            "almoco_x": round(x_a / mm, 1),
            "janta_x": round(x_j / mm, 1),
        }

    config = {
        "pagina": {
            "largura_mm": round(largura_mm, 1),
            "altura_mm": round(altura_mm, 1),
            "orientacao": "portrait",
        },
        "scan": {
            "dpi": 200,
        },
        "marcadores": {
            "tamanho_mm": round(marcador_mm, 1),
            "margem_mm": round(margem_mm, 1),
            "posicoes": marcadores,
        },
        "circulos": {
            "raio_mm": round(CIRCULO_RAIO / mm, 1),
            "borda_pt": CIRCULO_BORDA,
        },
        "layout": {
            "alunos_por_pagina": resultado_template["alunos_por_pagina"],
            "linha_altura_mm": round(LINHA_ALTURA / mm, 1),
            "primeira_linha_y_mm": round(
                altura_mm - resultado_template["primeiro_circulo_y_reportlab"] / mm,
                1
            ),
            "dias": dias,
            "circulos_por_dia": circulos_por_dia,
        },
        "restaurante": info["restaurante"],
    }

    with open(arquivo_config, "w", encoding="utf-8") as f:
        yaml.dump(config, f, default_flow_style=False, allow_unicode=True, sort_keys=False)

    print(f"Config gerado: {arquivo_config}")


# --- EXECUÇÃO ---

def main():
    if len(sys.argv) < 3:
        print("Uso: python gerar_template.py <planilha.xlsx> <nome_da_aba>")
        print()
        print("Exemplo:")
        print('  python gerar_template.py alunos.xlsx "CANELA IMPRESSÃO"')
        print()

        # Se só passou o arquivo, lista as abas disponíveis
        if len(sys.argv) == 2:
            wb = load_workbook(sys.argv[1], read_only=True)
            print(f"Abas disponíveis em {sys.argv[1]}:")
            for nome in wb.sheetnames:
                print(f"  - {nome}")
            wb.close()

        sys.exit(1)

    caminho_xlsx = sys.argv[1]
    nome_aba = sys.argv[2]

    # Lê os dados da planilha
    info = ler_planilha(caminho_xlsx, nome_aba)
    print(f"Lidos {len(info['alunos'])} alunos da aba '{nome_aba}'")

    # Determina os dias baseado na aba
    if nome_aba in CONFIG_ABAS:
        dias = CONFIG_ABAS[nome_aba]
    else:
        print(f"Aviso: aba '{nome_aba}' não tem configuração de dias definida.")
        print(f"Usando dias padrão (Seg-Sex).")
        dias = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta"]

    # Nome do arquivo de saída baseado no restaurante
    nome_restaurante = info["restaurante"].lower().replace(" ", "_") or "template"
    arquivo_saida = f"template_{nome_restaurante}.pdf"
    arquivo_config = f"config_{nome_restaurante}.yaml"

    resultado = gerar_template(info, dias, arquivo_saida)
    gerar_config(info, dias, resultado, arquivo_config)

    print(f"\nPosições dos círculos (em mm):")
    for i, dia in enumerate(dias):
        x_a, x_j = resultado["posicoes_circulos"][i]
        print(f"  {dia}: A = {x_a/mm:.1f}mm, J = {x_j/mm:.1f}mm")


if __name__ == "__main__":
    main()